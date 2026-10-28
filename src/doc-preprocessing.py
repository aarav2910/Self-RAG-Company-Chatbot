import hashlib
import os
from pathlib import Path

from dotenv import load_dotenv
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openai import OpenAI
from pypdf import PdfReader
from docx import Document as WordDocument

EMBEDDING_MODEL = "text-embedding-3-small"
PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DOCS_DIR = PROJECT_ROOT / "docs"
CHROMA_DIR = PROJECT_ROOT / "chroma_db"
COLLECTION_NAME = "document_chunks"
SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def create_clients():
    try:
        import chromadb
    except ImportError as exc:
        raise RuntimeError(
            "ChromaDB is not installed. Install it with: pip install chromadb"
        ) from exc

    load_dotenv(PROJECT_ROOT / ".env")
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is missing from the environment or .env file")

    openai_client = OpenAI(api_key=api_key)
    chroma_client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    collection = chroma_client.get_or_create_collection(name=COLLECTION_NAME)
    return openai_client, collection


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return [
            Document(
                page_content=page.extract_text() or "",
                metadata={"source": str(path), "page_number": page_number},
            )
            for page_number, page in enumerate(reader.pages, start=1)
        ]
    if suffix in {".txt", ".md"}:
        return [
            Document(
                page_content=path.read_text(encoding="utf-8"),
                metadata={"source": str(path), "page_number": 1},
            )
        ]
    if suffix == ".docx":
        doc = WordDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [
            Document(
                page_content=text,
                metadata={"source": str(path), "page_number": 1},
            )
        ]
    raise ValueError(f"Unsupported file type: {suffix}")


def load_all_documents(folder_path=DEFAULT_DOCS_DIR):
    folder = Path(folder_path)
    if not folder.is_dir():
        raise FileNotFoundError(f"Documents directory does not exist: {folder}")

    documents = []

    for file_path in sorted(folder.rglob("*")):
        if not file_path.is_file():
            continue

        if file_path.suffix.lower() not in SUPPORTED:
            continue

        try:
            print(f"Loading: {file_path}")

            documents.append({
                "filename": file_path.name,
                "document_id": hashlib.sha256(
                    file_path.read_bytes()
                ).hexdigest(),
                "pages": load_file(file_path),
            })

        except Exception as e:
            print(f"Failed to load {file_path}: {e}")

    return documents

def store_all_documents(documents, client, collection):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1300,
        chunk_overlap=200,
        add_start_index=True,
    )
    batch_size = 32

    for doc in documents:
        print(f"Processing: {doc['filename']}")

        chunks = splitter.split_documents(doc["pages"])

        ids = []
        texts = []
        metadatas = []

        for index, chunk in enumerate(chunks):
            content = chunk.page_content.strip()

            if not content:
                continue

            chunk_id = f"{doc['document_id']}_{index}"

            ids.append(chunk_id)
            texts.append(content)
            metadatas.append({
                "document_id": doc["document_id"],
                "filename": doc["filename"],
                "chunk_index": index,
                "page_number": chunk.metadata["page_number"],
                "start_index": chunk.metadata.get("start_index", 0),
            })

        for start in range(0, len(texts), batch_size):
            batch_ids = ids[start:start + batch_size]
            batch_texts = texts[start:start + batch_size]
            batch_metadatas = metadatas[start:start + batch_size]

            response = client.embeddings.create(
                model=EMBEDDING_MODEL,
                input=batch_texts,
            )

            vectors = [
                item.embedding
                for item in sorted(
                    response.data,
                    key=lambda item: item.index
                )
            ]

            collection.upsert(
                ids=batch_ids,
                documents=batch_texts,
                metadatas=batch_metadatas,
                embeddings=vectors,
            )

        print(f"Stored {len(texts)} chunks")


if __name__ == "__main__":
    client, collection = create_clients()
    documents = load_all_documents()
    print(f"Loaded {len(documents)} documents")

    if documents:
        store_all_documents(documents, client, collection)
    print("All documents processed!")
